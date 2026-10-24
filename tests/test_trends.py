import copy
import zipfile

import numpy as np
import pytest
from matplotlib.figure import Figure
from openpyxl import load_workbook
from PIL import Image

from plstudio.data import load_project, save_project
from plstudio.exporting import export_excel, export_trend_excel
from plstudio.fitting import fit_spectrum
from plstudio.plotting import DEFAULT_STYLE
from plstudio.trends import (DEFAULT_TRENDS, draw_trends, export_trend_set,
                             export_trends, prepare_trends)
from tests.test_core import synthetic


@pytest.fixture
def samples():
    first = synthetic(center=804)
    second = synthetic(center=807, amplitude=18000)
    first.condition, first.name = "Control", "Sample A"
    second.condition, second.name = "Treatment", "Sample B"
    for s in (first, second):
        s.result = fit_spectrum(s)
    return [first, second]


def test_trend_values_are_from_selected_peak_and_saved_result(samples):
    prepared = prepare_trends(samples)
    assert len(prepared["rows"]) == 8
    assert prepared["categories"] == ["Control", "Treatment"]
    first = prepared["rows"][0]
    assert first["metric"] == "maximum_position"
    assert first["value"] == samples[0].result["peaks"][0]["maximum_position"]
    assert first["stderr"] is None
    # Model parameters remain accessible to explicit numerical callers.
    center = prepare_trends(samples, {"view": "center"})["rows"][0]
    assert center["value"] == samples[0].result["peaks"][0]["center"]
    assert center["stderr"] == samples[0].result["peaks"][0]["center_stderr"]
    assert first["unit"] == "nm"
    # Working settings may change; a trend must continue to use the saved fit's units.
    samples[0].config.domain = "eV"
    assert prepare_trends(samples)["rows"][0]["unit"] == "nm"
    second_peak = copy.deepcopy(samples[0].result["peaks"][0])
    second_peak.update(peak=2, center=850, maximum_position=850, fwhm=5)
    samples[0].result["peaks"].append(second_peak)
    only_peak2 = prepare_trends(samples, {"peak": 2})
    assert only_peak2["rows"][0]["value"] == 850
    assert only_peak2["rows"][4]["value"] is None
    assert "沒有 Peak 2" in only_peak2["rows"][4]["status"]


def test_category_numeric_axes_duplicates_and_shared_selection(samples):
    samples[1].condition = "Control"
    p = prepare_trends(samples)
    assert p["categories"] == ["Control"]
    assert len([r for r in p["rows"] if r["metric"] == "maximum_position"]) == 2
    assert [r["x"] for r in p["rows"] if r["metric"] == "maximum_position"] == [0, 0]
    samples[0].plot = False
    assert prepare_trends(samples)["spectrum_count"] == 1
    samples[0].plot = True
    with pytest.raises(ValueError, match="有限數字"):
        prepare_trends(samples, {"x_mode": "numeric"})
    samples[0].condition, samples[1].condition = "100", "20"
    prepared = prepare_trends(samples, {"x_mode": "auto"})
    assert prepared["x_mode"] == "numeric"
    figure = Figure()
    axes, _ = draw_trends(figure, samples, {"x_mode": "numeric", "view": "center"})
    assert list(axes[0].lines[0].get_xdata()) == [20, 100]
    samples[0].condition = "nan"
    with pytest.raises(ValueError, match="有限數字"):
        prepare_trends(samples, {"x_mode": "numeric"})


def test_failed_missing_nonfinite_results_and_compatible_units(samples):
    samples[1].result["success"] = False
    prepared = prepare_trends(samples)
    assert sum(r["valid"] for r in prepared["rows"]) == 4
    assert sum(r["valid"] for r in prepare_trends(samples, {"include_failed": True})["rows"]) == 8
    samples[1].result["success"] = True
    samples[1].result["peaks"][0]["fwhm"] = None
    prepared = prepare_trends(samples)
    assert any(r["value"] is None and r["metric"] == "fwhm" for r in prepared["rows"])
    samples[1].result["config"]["domain"] = "eV"
    with pytest.raises(ValueError, match="不一致"):
        prepare_trends(samples)
    # A normalized R² is dimensionless and remains available as a diagnostic.
    assert len(prepare_trends(samples, {"view": "r2"})["rows"]) == 2
    samples[1].result = None
    assert "尚未擬合" in prepare_trends(samples)["rows"][4]["status"]


def test_interval_warning_and_no_fabricated_fwhm_uncertainty(samples):
    samples[1].result["config"]["xmin"] = 750
    prepared = prepare_trends(samples)
    assert any("積分區間" in m for m in prepared["messages"])
    assert all(r["stderr"] is None for r in prepared["rows"] if r["metric"] == "fwhm")
    samples[1].result["config"]["normalization"] = "max"
    # Widths can be compared under a pure scalar normalization, areas cannot.
    prepare_trends(samples, {"view": "fwhm"})
    with pytest.raises(ValueError, match="不一致"):
        prepare_trends(samples, {"view": "area_in_range"})


def test_project_backward_compatibility_condition_and_settings(samples, tmp_path):
    samples[0].plot = samples[0].trend_include = False
    style = DEFAULT_STYLE | {"trends": DEFAULT_TRENDS | {"xlabel": "Power (mW)", "view": "fwhm"}}
    path = tmp_path/"condition.plproj"
    save_project(path, samples, style)
    restored, settings = load_project(path)
    assert restored[0].condition == "Control"
    assert not restored[0].trend_include
    assert settings["trends"]["view"] == "fwhm"
    # Old schema-1 projects did not contain condition fields.
    import json
    with zipfile.ZipFile(path) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    manifest = json.loads(members["project.json"])
    for record in manifest["spectra"]:
        record.pop("condition")
        record.pop("trend_include")
    members["project.json"] = json.dumps(manifest).encode()
    old_path = tmp_path/"old.plproj"
    with zipfile.ZipFile(old_path, "w") as archive:
        for name, content in members.items():
            archive.writestr(name, content)
    restored, _ = load_project(old_path)
    assert restored[0].condition == ""
    assert restored[0].trend_include == restored[0].plot


def test_trend_image_and_excel_exports(samples, tmp_path):
    style = DEFAULT_TRENDS | {"width_cm": 15.24, "height_cm": 10.16, "dpi": 100}
    for extension in ("png", "jpg", "pdf", "svg"):
        export_trends(tmp_path/f"overview.{extension}", samples, style)
        assert (tmp_path/f"overview.{extension}").stat().st_size > 1000
    with Image.open(tmp_path/"overview.png") as im:
        assert im.size == (600, 400)
    paths, skipped = export_trend_set(tmp_path, samples, style)
    assert len(paths) == 4 and not skipped
    export_trend_excel(tmp_path/"trends.xlsx", samples, style)
    book = load_workbook(tmp_path/"trends.xlsx")
    sheet = book["Trend data"]
    headers = [cell.value for cell in sheet[1]]
    assert sheet.max_row == 9
    intensity_rows = [r for r in prepare_trends(samples)["rows"] if r["metric"] == "height"]
    assert intensity_rows[0]["value"] == samples[0].result["peaks"][0]["height"]
    assert intensity_rows[0]["unit"] == "Counts"
    assert "Peak intensity" in intensity_rows[0]["metric_label"]
    assert isinstance(sheet.cell(2, headers.index("value")+1).value, (int, float))
    assert sheet.cell(2, headers.index("condition")+1).value == "Control"
    samples[0].condition = "=1+1"
    export_excel(tmp_path/"full.xlsx", samples)
    sheet = load_workbook(tmp_path/"full.xlsx")["Summary"]
    headers = [cell.value for cell in sheet[1]]
    cell = sheet.cell(2, headers.index("condition")+1)
    assert cell.value == "=1+1" and cell.data_type == "s"
    peak_sheet = load_workbook(tmp_path/"full.xlsx")["Peak parameters"]
    peak_headers = [cell.value for cell in peak_sheet[1]]
    assert peak_sheet.cell(2, peak_headers.index("peak_intensity")+1).value == pytest.approx(
        samples[0].result["peaks"][0]["height"])


def test_no_metrics_or_no_eligible_data(samples, tmp_path):
    with pytest.raises(ValueError, match="至少勾選"):
        prepare_trends(samples, {"metrics": []})
    for s in samples:
        s.plot = s.trend_include = False
    with pytest.raises(ValueError, match="沒有可繪製"):
        export_trends(tmp_path/"empty.png", samples)
