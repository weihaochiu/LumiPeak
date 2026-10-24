import copy
import json
import threading
import zipfile
from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook

from plstudio.data import (HC, FitConfig, import_path, load_project, parse_spectrum,
                           save_project, transformed)
from plstudio.exporting import export_excel, export_origin_bundle
from plstudio.fitting import (MODELS, FitCancelled, build_model, fit_joint,
                              fit_spectrum, parameter_config)
from plstudio.plotting import DEFAULT_STYLE, export_figure


def synthetic(center=805.2, sigma=9.3, amplitude=25000, baseline=700):
    x = np.linspace(730, 880, 351)
    rng = np.random.default_rng(42)
    y = baseline+amplitude/(sigma*np.sqrt(2*np.pi))*np.exp(-.5*((x-center)/sigma)**2)
    y += rng.normal(0, 2, len(x))
    raw = "\n".join(f"{xx:.10f}\t{yy:.10f}" for xx, yy in zip(x, y))
    return parse_spectrum((raw+"\nData Type: Counts\t--\nTemperature (C): -50\t--").encode(), "sample.txt")


def test_parse_metadata_and_duplicate(tmp_path):
    s = synthetic()
    assert len(s.x) == 351
    assert ["Temperature (C)", "-50"] in s.metadata
    with zipfile.ZipFile(tmp_path/"sample.zip", "w") as z:
        z.writestr("one.txt", s.raw)
        z.writestr("different_name.txt", s.raw)
        z.writestr("bad.txt", b"no data")
    ss, errors = import_path(tmp_path/"sample.zip")
    assert len(ss) == 2 and len(errors) == 1
    assert ss[0].digest == ss[1].digest


def test_transform_preserves_density_integral():
    s = synthetic()
    c = FitConfig(domain="eV", jacobian=True)
    x, y = transformed(s, c)
    assert np.all(np.diff(x) > 0)
    assert np.trapezoid(y, x) == pytest.approx(np.trapezoid(s.y, s.x), rel=2e-5)
    c.jacobian = False
    _, yy = transformed(s, c)
    assert np.allclose(yy, s.y[::-1])


def test_gaussian_recovery_and_fixed_constraint():
    s = synthetic()
    result = fit_spectrum(s)
    p = result["peaks"][0]
    assert result["success"]
    assert p["center"] == pytest.approx(805.2, abs=.02)
    assert p["fwhm"] == pytest.approx(2.354820045*9.3, rel=.002)
    assert p["amplitude"] == pytest.approx(25000, rel=.002)
    assert result["metrics"]["rmse"] < 3
    assert all(p["stderr"] is not None for p in result["parameters"])
    _, params, _, _ = build_model(s, s.config)
    config = copy.deepcopy(s.config)
    config.parameters = parameter_config(params)
    config.parameters["p1_center"].update(value=804.0, vary=False)
    result = fit_spectrum(s, config)
    assert result["peaks"][0]["center"] == 804.0
    config.parameters["p1_sigma"]["value"] = -1
    with pytest.raises(ValueError, match="初始值超出"):
        fit_spectrum(s, config)


@pytest.mark.parametrize("name", list(MODELS))
def test_model_can_fit_own_spectrum(name):
    s = synthetic()
    c = FitConfig(peaks=[name], starts=1)
    model, params, x, _ = build_model(s, c)
    params["p1_center"].value = 804
    params["p1_amplitude"].value = 30000
    params["p1_sigma"].value = 8
    if "p1_gamma" in params:
        params["p1_gamma"].value = 4
    if "p1_skew" in params:
        params["p1_skew"].value = 1.1
    s.y = model.eval(params, x=x)
    result = fit_spectrum(s, c)
    assert result["success"]
    assert result["metrics"]["rmse"] < .05


def test_expression_constraints_and_all_fixed():
    s = synthetic()
    c = FitConfig()
    _, params, _, _ = build_model(s, c)
    c.parameters = parameter_config(params)
    c.parameters["p1_center"]["expr"] = "800 + p1_sigma / 2"
    result = fit_spectrum(s, c)
    pp = {p["parameter"]: p["value"] for p in result["parameters"]}
    assert pp["p1_center"] == pytest.approx(800+pp["p1_sigma"]/2)
    c.parameters["p1_center"]["expr"] = ""
    for p in c.parameters.values():
        p["vary"] = False
    result = fit_spectrum(s, c)
    assert result["success"] and result["metrics"]["nfree"] == 0


def test_joint_shares_parameters_and_uses_joint_scores():
    s1, s2 = synthetic(amplitude=25000), synthetic(amplitude=12000)
    configs = [FitConfig(shared=["p1_center", "p1_sigma"], starts=1) for _ in range(2)]
    results = fit_joint([s1, s2], configs, configs[0].shared)
    assert results[0]["peaks"][0]["center"] == results[1]["peaks"][0]["center"]
    assert results[0]["peaks"][0]["center_stderr"] == results[1]["peaks"][0]["center_stderr"]
    assert results[0]["joint"]["id"] == results[1]["joint"]["id"]
    assert results[0]["joint"]["metrics"]["nfree"] == 6
    assert results[0]["metrics"]["aic"] is None


def test_cancel_and_masks():
    s = synthetic()
    event = threading.Event()
    event.set()
    with pytest.raises(FitCancelled):
        fit_spectrum(s, cancel=event)
    c = FitConfig(excluded=[[800, 802]])
    result = fit_spectrum(s, c)
    assert not any(800 <= x <= 802 for x in result["x"])
    assert result["metrics"]["measured_area_background_subtracted"] is None


def test_project_roundtrip_integrity_and_literal_excel(tmp_path):
    s = synthetic()
    s.result = fit_spectrum(s)
    s.name = '=HYPERLINK("http://example.com")'
    style = copy.deepcopy(DEFAULT_STYLE)
    save_project(tmp_path/"test.plproj", [s], style)
    loaded, new_style = load_project(tmp_path/"test.plproj")
    assert loaded[0].raw == s.raw
    assert np.array_equal(loaded[0].x, s.x)
    assert loaded[0].result == s.result
    assert new_style == style
    export_excel(tmp_path/"results.xlsx", loaded)
    book = load_workbook(tmp_path/"results.xlsx")
    assert book["Summary"]["B2"].value == s.name
    assert book["Summary"]["B2"].data_type == "s"
    headers = [c.value for c in book["Peak parameters"][1]]
    assert isinstance(book["Peak parameters"].cell(2, headers.index("center")+1).value, float)
    export_origin_bundle(tmp_path/"origin.zip", loaded)
    with zipfile.ZipFile(tmp_path/"origin.zip") as z:
        assert any(n.endswith("_fit.csv") for n in z.namelist())
        assert "import_origin.py" in z.namelist()
        compile(z.read("import_origin.py"), "import_origin.py", "exec")


def test_figure_exports(tmp_path):
    from PIL import Image
    s = synthetic()
    s.result = fit_spectrum(s)
    style = DEFAULT_STYLE | {"width_cm": 15.24, "height_cm": 10.16, "dpi": 100,
                             "residual": True, "components": True}
    for ext in ("png", "jpg", "pdf", "svg"):
        export_figure(tmp_path/f"figure.{ext}", [s], style)
        assert (tmp_path/f"figure.{ext}").stat().st_size > 1000
    with Image.open(tmp_path/"figure.png") as im:
        assert im.size == (600, 400)


@pytest.mark.skipif(not Path(r"D:\檔案交換資料夾\PL.zip").exists(), reason="Reference archive unavailable")
def test_user_raw_archive(tmp_path):
    spectra, errors = import_path(r"D:\檔案交換資料夾\PL.zip")
    assert len(spectra) == 6 and not errors
    assert all(len(s.x) == 1024 for s in spectra)
    by_name = {s.filename: s for s in spectra}
    assert by_name["C0.txt"].digest == by_name["C0-1.txt"].digest
    for spectrum in spectra:
        spectrum.result = fit_spectrum(spectrum)
        assert spectrum.result["success"]
        assert 795 < spectrum.result["peaks"][0]["center"] < 820
    export_excel(tmp_path/"raw_results.xlsx", spectra)
    save_project(tmp_path/"raw.plproj", spectra, DEFAULT_STYLE)
    restored, _ = load_project(tmp_path/"raw.plproj")
    assert len(restored) == 6
