from __future__ import annotations

from .i18n import tr

import csv
import io
import json
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .data import FitConfig, clean_json, intensity_label


def _identity(spectrum):
    return {"spectrum_id": spectrum.id, "display_name": spectrum.name,
            "original_filename": spectrum.filename, "condition": spectrum.condition.strip() or spectrum.name}


def result_rows(spectra):
    summary, peaks, parameters, settings, metadata, history, joints = [], [], [], [], [], [], []
    seen_joint = set()
    for spectrum in spectra:
        identity = _identity(spectrum)
        result = spectrum.result
        if result:
            config = FitConfig(**result["config"])
            summary.append(identity | {"fit_id": result["id"], "success": result["success"],
                           "domain": config.domain, "intensity": intensity_label(config),
                           "normalization": config.normalization, "jacobian": config.jacobian,
                           "model": "+".join(config.peaks), "background": config.background,
                           "warnings": "; ".join(result["warnings"]),
                           "joint_fit_id": result["joint"]["id"] if result["joint"] else "",
                           "timestamp_UTC": result["timestamp"],
                           "peak_1_intensity": result["peaks"][0]["height"]} | result["metrics"])
            peaks.extend(identity | {"fit_id": result["id"], "x_unit": config.domain,
                          "intensity": intensity_label(config), "normalization": config.normalization,
                          "jacobian": config.jacobian,
                          "area_definition": "model integral over full fit limits, including masked intervals",
                          "peak_intensity": p["height"],
                          "peak_intensity_definition": "maximum of fitted peak component, excluding background"} | p
                         for p in result["peaks"])
            parameters.extend(identity | {"fit_id": result["id"], "x_unit": config.domain} | p for p in result["parameters"])
            settings.append(identity | {"config_json": json.dumps(result["config"], ensure_ascii=False),
                                         "error_assumption": result["error_assumption"],
                                         "message": result["message"]})
            joint = result["joint"]
            if joint and joint["id"] not in seen_joint:
                joints.append({"joint_fit_id": joint["id"], "shared": ", ".join(joint["shared"]),
                               "members": ", ".join(joint["members"]), "note": joint["note"]} | joint["metrics"])
                seen_joint.add(joint["id"])
        else:
            summary.append(identity | {"success": "Not fitted"})
        metadata.extend(identity | {"field": k, "value": v} for k, v in spectrum.metadata)
        metadata.append(identity | {"field": "User notes", "value": spectrum.notes})
        history.extend(identity | {"fit_id": r["id"], "timestamp_UTC": r["timestamp"],
                                    "model": "+".join(r["config"]["peaks"]),
                                    "success": r["success"], "config_json": json.dumps(r["config"], ensure_ascii=False)}
                       for r in spectrum.history)
    return {"Summary": summary, "Peak parameters": peaks, "All parameters": parameters,
            "Fit settings": settings, "Instrument metadata": metadata,
            "Fit history": history, "Joint fit metrics": joints}


def export_excel(path, spectra):
    if not spectra:
        raise ValueError(tr("沒有可匯出的資料。"))
    _export_tables(path, result_rows(spectra))


def export_trend_excel(path, spectra, settings):
    from .trends import prepare_trends
    prepared = prepare_trends(spectra, settings, overview=True)
    if not prepared["rows"]:
        raise ValueError(tr("沒有勾選的趨勢資料。"))
    rows = []
    for row in prepared["rows"]:
        record = {k: v for k, v in row.items() if k != "fit_config"}
        record["fit_config_json"] = json.dumps(row["fit_config"], ensure_ascii=False)
        rows.append(record)
    _export_tables(path, {"Trend data": rows,
                         "Trend settings": [{"setting": k, "value": json.dumps(v, ensure_ascii=False)}
                                            for k, v in settings.items()],
                         "Notices": [{"notice": message} for message in prepared["messages"]]})


def _export_tables(path, tables):
    book = Workbook()
    book.remove(book.active)
    for title, rows in tables.items():
        sheet = book.create_sheet(title)
        headers = list(dict.fromkeys(key for row in rows for key in row))
        if not headers:
            headers = ["No data"]
        sheet.append(headers)
        for row in rows:
            sheet.append([clean_json(row.get(k)) for k in headers])
        # User-provided names and metadata remain literal text, never Excel formulas.
        for cells in sheet:
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                elif isinstance(cell.value, float):
                    cell.number_format = "0.000000E+00"
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="173B5A")
            cell.alignment = Alignment(vertical="center")
        sheet.freeze_panes = "D2"
        sheet.auto_filter.ref = sheet.dimensions
        for column, header in enumerate(headers, 1):
            width = max([len(str(header)), *[len(str(row.get(header, ""))) for row in rows[:100]]])
            sheet.column_dimensions[get_column_letter(column)].width = min(55, max(14, width+2))
    book.save(path)


def _csv(headers, rows):
    out = io.StringIO(newline="")
    writer = csv.writer(out)
    writer.writerow(headers)
    writer.writerows(rows)
    return out.getvalue().encode("utf-8-sig")


ORIGIN_SCRIPT = '''"""Run using a Python environment with originpro and licensed Origin installed.
Extract this ZIP to a folder, then run: python import_origin.py
Worksheets and editable line graphs are created. Python figure styling is not transferred.
"""
from pathlib import Path
import csv
import originpro as op

folder = Path(__file__).resolve().parent
if list(op.pages()):
    raise RuntimeError("Origin already contains pages. Save your work, then use File > New Project, "
                       "remove its empty windows and run this script again.")
op.set_show(True)
for file in sorted(folder.glob("*_fit.csv")):
    with file.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))
    sheet = op.new_sheet("w", lname=file.stem)
    headers = rows[0]
    for j, header in enumerate(headers):
        sheet.from_list(j, [float(row[j]) for row in rows[1:]], lname=header)
    graph = op.new_graph(template="line")
    layer = graph[0]
    for column in (1, 2):
        layer.add_plot(sheet, column, 0, type="l")
    layer.rescale()
for file in sorted(folder.glob("*_raw.csv")):
    with file.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))
    sheet = op.new_sheet("w", lname=file.stem)
    for j, header in enumerate(rows[0]):
        sheet.from_list(j, [float(row[j]) for row in rows[1:]], lname=header)
op.save(str(folder / "PL_results.opju"))
print("Saved", folder / "PL_results.opju")
'''


def export_origin_bundle(path, spectra):
    if not spectra:
        raise ValueError(tr("沒有可匯出的資料。"))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("import_origin.py", ORIGIN_SCRIPT.encode("utf-8"))
        archive.writestr("README.txt", (
            "PL Studio Origin data bundle\nExtract all files before running import_origin.py.\n"
            "Install originpro in an external Python environment with licensed Origin on Windows.\n"
            "CSV files can also be imported manually. All curves remain numeric and editable.\n"
            "Plot style is not transferred; fit CSV contains ONLY included fitting points.\n"
            "Raw CSV preserves the full original acquisition.\n").encode("utf-8"))
        archive.writestr("results.json", json.dumps(clean_json(result_rows(spectra)), ensure_ascii=False))
        for i, spectrum in enumerate(spectra, 1):
            stem = f"{i:03d}_{spectrum.id[:8]}"
            archive.writestr(stem+"_raw.csv", _csv(["Wavelength_nm", "Counts"], zip(spectrum.x, spectrum.y)))
            if spectrum.result:
                result = spectrum.result
                keys = list(result["components"])
                headers = ["X_"+result["config"]["domain"], "Data", "Fit", "Residual", *keys]
                columns = [result["x"], result["y"], result["fitted"], result["residual"],
                           *[result["components"][k] for k in keys]]
                archive.writestr(stem+"_fit.csv", _csv(headers, zip(*columns)))
        archive.writestr("spectrum_names.json", json.dumps(
            {f"{i:03d}_{s.id[:8]}": _identity(s) for i, s in enumerate(spectra, 1)}, ensure_ascii=False))


def export_opju(path, spectra):
    try:
        import originpro as op
    except ImportError as exc:
        raise RuntimeError(tr("此環境尚未安裝 originpro。可先匯出 Origin 資料包；\n"
                           "直接建立 .opju 需要安裝並授權 Origin，以及相容的 originpro。")) from exc
    # External originpro launches its own instance by default; never attach to a user's session.
    # Refuse to replace existing pages even if the environment is configured differently.
    if list(op.pages()):
        raise RuntimeError(tr("此 Origin 連線已有工作頁。請先儲存，建立空白專案並關閉其中的空白視窗，\n"
                           "或使用 Origin 資料包匯入，以免覆蓋現有工作。"))
    for i, spectrum in enumerate(spectra, 1):
        sheet = op.new_sheet("w", lname=f"{i}_{spectrum.name}")
        result = spectrum.result
        if result:
            columns = {"X_"+result["config"]["domain"]: result["x"], "Data": result["y"],
                       "Fit": result["fitted"], "Residual": result["residual"], **result["components"]}
        else:
            columns = {"Wavelength_nm": spectrum.x.tolist(), "Counts": spectrum.y.tolist()}
        for j, (key, values) in enumerate(columns.items()):
            sheet.from_list(j, values, lname=key)
        graph = op.new_graph(template="line")
        layer = graph[0]
        layer.add_plot(sheet, 1, 0, type="l")
        if result:
            layer.add_plot(sheet, 2, 0, type="l")
        layer.rescale()
    if not op.save(str(Path(path).resolve())):
        raise RuntimeError(tr("Origin 未成功儲存專案。"))
